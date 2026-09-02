"""Generates docs/Ledger-Oracle-Pitch.pptx -- 6 slides, sourced from video-brief.md.
Run: python docs/build_pitch_deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK = RGBColor(0x1c, 0x18, 0x10)
PAPER = RGBColor(0xf4, 0xef, 0xe4)
PANEL = RGBColor(0xfb, 0xf9, 0xf3)
LINE = RGBColor(0xd8, 0xcf, 0xb8)
ACCENT = RGBColor(0xa3, 0x7c, 0x1f)
DIM = RGBColor(0x5c, 0x53, 0x42)
PASS = RGBColor(0x3a, 0x7a, 0x4c)
BLOCK = RGBColor(0x9c, 0x38, 0x30)

SERIF = "Georgia"
SANS = "Calibri"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    return s


def box(s, l, t, w, h):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, size, color=INK, bold=False, font=SANS, first=False, align=PP_ALIGN.LEFT, space_after=6):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    return p


def rect(s, l, t, w, h, fill=PANEL, line=LINE):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def eyebrow(s, text, num):
    tf = box(s, 0.6, 0.35, 8, 0.4)
    para(tf, f"{num}  ·  LEDGER ORACLE  ·  RAZORPAY AI BUILDATHON 2026", 11, ACCENT, True, MONO, first=True)


def footer_rule(s):
    from pptx.enum.shapes import MSO_SHAPE
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.05), Inches(12.13), Pt(1.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); ln.shadow.inherit = False


# ---------- Slide 1: Title ----------
s = add_slide()
tf = box(s, 1, 2.2, 11.3, 1.6)
para(tf, "The claim isn't the fact.", 44, INK, False, SERIF, first=True)
para(tf, "The ledger is.", 44, ACCENT, True, SERIF)
tf2 = box(s, 1, 3.9, 10.5, 1.2)
para(tf2, "Ledger Oracle — an AI agent that checks refund claims against the real payments\nledger before money moves.", 18, DIM, False, SANS, first=True)
tf3 = box(s, 1, 6.3, 10.5, 0.6)
para(tf3, "Razorpay AI Buildathon 2026  ·  Track 02 — AI Risk Manager", 13, DIM, False, MONO, first=True)

# ---------- Slide 2: Problem ----------
s = add_slide()
eyebrow(s, "01 — THE PROBLEM", "")
tf = box(s, 0.6, 0.95, 11, 0.9)
para(tf, "Every guardrail checks if a claim is allowed.\nNone of them check if it's true.", 30, INK, True, SERIF, first=True)

rect(s, 0.6, 2.2, 5.7, 4.3)
tf = box(s, 0.85, 2.4, 5.2, 4.0)
para(tf, "CUSTOMER", 11, ACCENT, True, MONO, first=True)
para(tf, '"I was charged twice for order #4471.\nPlease refund one of them."', 16, INK, False, SANS, space_after=14)
para(tf, "WHAT A CONFIG-ONLY GUARDRAIL SEES", 11, BLOCK, True, MONO)
para(tf, "Amount ₹2,499 — inside limit. Scope, consent, PII —\nall clear. Approved.\n\nThere was one charge, not two. The refund pays out\nanyway, because nothing upstream ever looked at the\nledger.", 14.5, INK, False, SANS)

rect(s, 6.6, 2.2, 6.13, 4.3, fill=PAPER, line=LINE)
tf = box(s, 6.85, 2.45, 5.6, 3.9)
para(tf, "THE GAP", 11, ACCENT, True, MONO, first=True)
para(tf, "claim  →  agent decides  →  config guardrails  →  money moves", 13.5, DIM, False, MONO, space_after=16)
para(tf, "Nothing in that chain ever asks:\n\"does the ledger agree with this story?\"", 16, INK, True, SERIF, space_after=16)
para(tf, "Four major Indian PSPs have publicly written about this\nexact scam. Their own advice: “verify the UTR against\nyour bank statement” — precisely the check nothing in\nthe stack automates.", 13.5, DIM)
footer_rule(s)

# ---------- Slide 3: What we built / architecture ----------
s = add_slide()
eyebrow(s, "02 — HOW IT DECIDES", "")
tf = box(s, 0.6, 0.95, 11.5, 0.9)
para(tf, "A bounded agent investigates. A deterministic\nfunction decides. The model never authorizes money.", 26, INK, True, SERIF, first=True)

cols = [
    ("STEP 1 — INVESTIGATION AGENT", "agent.py", "Bounded LLM. Max 5 tool calls, 2 retries\neach, 30s wall clock. Out of budget →\nescalate. Never hangs, never guesses.", INK),
    ("5 READ-ONLY TOOLS", "tools.py", "get_payment_by_utr, get_order_payments,\nfind_duplicate_captures, check_refund_history,\ncheck_payment_status. mutates=False on every\none. No issue_refund tool exists — on purpose.", INK),
    ("STEP 2 — POLICY ENGINE", "policy.py — decide()", '"DECIDES. No network. No model. No\nrandomness." Plain comparison against\nledger evidence. Returns pass / block /\nescalate.', ACCENT),
]
x = 0.6
for title, name, desc, accentcol in cols:
    rect(s, x, 2.15, 3.95, 3.0, fill=PANEL, line=LINE)
    tf = box(s, x + 0.25, 2.35, 3.5, 2.7)
    para(tf, title, 10.5, accentcol, True, MONO, first=True, space_after=4)
    para(tf, name, 14, INK, True, MONO, space_after=8)
    para(tf, desc, 12, DIM, False, SANS)
    x += 4.2

tf = box(s, 0.6, 5.4, 11.7, 1.4)
para(tf, "PASS", 14, PASS, True, MONO, first=True, space_after=2)
para(tf, "Claim matches ledger evidence within tolerance — refund authorized.", 12.5, DIM, False, SANS, space_after=8)
para(tf, "BLOCK", 14, BLOCK, True, MONO, space_after=2)
para(tf, "Claim contradicts the ledger — reused reference, amount mismatch, exceeds captured total.", 12.5, DIM, False, SANS, space_after=8)
para(tf, "ESCALATE", 14, ACCENT, True, MONO, space_after=2)
para(tf, "Evidence insufficient or in-flight — a human decides. The engine never guesses its way to a pass.", 12.5, DIM, False, SANS)
footer_rule(s)

# ---------- Slide 4: Safety — abuse guard + replay ----------
s = add_slide()
eyebrow(s, "03 — SAFETY, NOT JUST A LOOKUP", "")
tf = box(s, 0.6, 0.95, 11.5, 0.9)
para(tf, "Two guards close the loop a lookup alone leaves open.", 26, INK, True, SERIF, first=True)

rect(s, 0.6, 2.15, 5.85, 4.3, fill=PANEL, line=LINE)
tf = box(s, 0.85, 2.35, 5.35, 4.0)
para(tf, "REPEAT-CLAIM ABUSE GUARD", 12, ACCENT, True, MONO, first=True, space_after=8)
para(tf, "3 prior claims on the same order that ended in a\ngenuine, terminal denial → the next automatic\nattempt blocks outright, before any tool call runs.", 14, INK, False, SANS, space_after=10)
para(tf, "Only counts denials the customer wasn't invited to\nretry — an in-flight capture, a down model, or a typo\nnever counts against an honest customer.", 13, DIM, False, SANS)

rect(s, 6.65, 2.15, 6.08, 4.3, fill=PANEL, line=LINE)
tf = box(s, 6.9, 2.35, 5.6, 4.0)
para(tf, "REPLAY GUARD", 12, ACCENT, True, MONO, first=True, space_after=8)
para(tf, "The moment a claim resolves to pass, its reference is\nimmediately consumed. Filing the exact same claim\nagain blocks with REF_ALREADY_CONSUMED.", 14, INK, False, SANS, space_after=10)
para(tf, "One refund per reference. Not infinite. Proven live in\nthe demo, not claimed in a slide.", 13, DIM, False, SANS)
footer_rule(s)

# ---------- Slide 5: Proof numbers ----------
s = add_slide()
eyebrow(s, "04 — PROOF, NOT PROMISES", "")
tf = box(s, 0.6, 0.95, 11.5, 0.9)
para(tf, "Numbers checked before writing them down.", 26, INK, True, SERIF, first=True)

stats = [
    ("0 / 1", "unsafe_pass_count on a dedicated adversarial\ncorpus — prompt injection, homoglyph & zero-\nwidth UTR disguises, contradictory amounts."),
    ("100%", "of the agent's tool allowlist is read-only by\nconstruction — the DB connection itself can't write."),
    ("138", "automated tests — several assert the architecture\nitself, e.g. decide() makes zero network calls."),
    ("HASH-CHAINED", "append-only audit log. Tamper one field and the\nverifier reports the exact broken event, live."),
]
x = 0.6
for n, d in stats:
    rect(s, x, 2.2, 2.85, 4.1, fill=PANEL, line=LINE)
    tf = box(s, x + 0.2, 2.45, 2.45, 3.7)
    para(tf, n, 22 if len(n) < 8 else 15, ACCENT, True, MONO, first=True, space_after=10)
    para(tf, d, 11.5, DIM, False, SANS)
    x += 3.0
footer_rule(s)

# ---------- Slide 6: Live demo / CTA ----------
s = add_slide()
eyebrow(s, "05 — SEE IT LIVE", "")
tf = box(s, 0.6, 0.95, 11.5, 0.9)
para(tf, "Every step below is a real write to a real database.", 26, INK, True, SERIF, first=True)

steps = [
    ("1", "Simulate a payment", "/pay — real row lands in captures, real UTR generated. Screen says “Payment failed” anyway — the exact gap this product catches."),
    ("2", "File the complaint", "One click hands the real UTR to the claim form. Agent investigates live, finds the match — pass, refund authorized."),
    ("3", "Replay the same claim", "Blocked. REF_ALREADY_CONSUMED. One refund, not infinite."),
    ("4", "Open the admin console", "Claim history, engine-vs-human agreement, audit-chain integrity — checkable live."),
]
y = 2.15
for n, t, d in steps:
    rect(s, 0.6, y, 0.55, 0.9, fill=ACCENT, line=ACCENT)
    tfn = box(s, 0.6, y + 0.1, 0.55, 0.7)
    para(tfn, n, 20, PAPER, True, MONO, first=True, align=PP_ALIGN.CENTER)
    tf = box(s, 1.35, y, 11.0, 1.05)
    p = para(tf, t, 15, INK, True, SANS, first=True, space_after=2)
    para(tf, d, 12.5, DIM, False, SANS)
    y += 1.15

tf = box(s, 0.6, 6.75, 11.5, 0.5)
para(tf, "github.com/abhinavteja123/Ledger-Oracle", 12, ACCENT, True, MONO, first=True)

prs.save("docs/Ledger-Oracle-Pitch.pptx")
print("saved docs/Ledger-Oracle-Pitch.pptx")
